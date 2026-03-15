from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def style_header(row) -> None:
    for cell in row:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--name", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    wb = Workbook()
    ws = wb.active
    ws.title = "实验设计"
    ws.append(["字段", "内容"])
    style_header(ws[1])
    for item in [
        "实验名称",
        "业务线",
        "实验目标",
        "核心假设",
        "实验组",
        "对照组",
        "分群方式",
        "主指标",
        "辅指标",
        "风险指标",
        "上线时间",
        "结束时间",
    ]:
        ws.append([item, ""])

    ws2 = wb.create_sheet("数据口径")
    ws2.append(["指标", "定义", "来源表", "字段", "备注"])
    style_header(ws2[1])

    ws3 = wb.create_sheet("复盘")
    ws3.append(["结论项", "内容"])
    style_header(ws3[1])
    for item in ["核心结果", "关键发现", "归因判断", "后续动作"]:
        ws3.append([item, ""])

    for sheet in wb.worksheets:
        sheet.column_dimensions["A"].width = 20
        sheet.column_dimensions["B"].width = 40
        sheet.freeze_panes = "A2"

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(f"Wrote workbook to {args.output}")


if __name__ == "__main__":
    main()
