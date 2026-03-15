from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


SECTIONS = [
    "1. 背景",
    "2. 目标",
    "3. 用户范围",
    "4. 用户路径",
    "5. 页面结构",
    "6. 模块说明",
    "7. 交互说明",
    "8. 视觉说明",
    "9. 规则与异常处理",
    "10. 埋点与数据口径",
    "11. 风险与依赖",
    "12. 排期建议",
]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--title", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(args.title)
    run.bold = True
    run.font.size = Pt(18)

    meta = doc.add_paragraph()
    meta.add_run("文档类型：").bold = True
    meta.add_run("微信活动 PRD")

    for section in SECTIONS:
        doc.add_heading(section, level=1)
        doc.add_paragraph("请在这里补充内容。")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(f"Wrote docx to {args.output}")


if __name__ == "__main__":
    main()
